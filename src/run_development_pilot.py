"""Execute the two-drug, development-only nested-CV/resource pilot.

Never opens CTRP response files. A read-only guard blocks them if future code
accidentally attempts access. Cohort-v1 frozen artifacts are never rewritten.
"""
from __future__ import annotations
import argparse,csv,hashlib,importlib.metadata,json,os,platform,sys,threading,time,warnings
from pathlib import Path
import numpy as np
import pandas as pd
import psutil
from openpyxl import load_workbook
from threadpoolctl import threadpool_limits,threadpool_info
from modeling.features import FoldFeatureBuilder
from modeling.nested_cv import make_grouped_splits,evaluate_outer_fold,metrics
from project_paths import REPOSITORY,original_input

ROOT=REPOSITORY
ACTIVE_RUN={}

def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(8*1024*1024),b''):h.update(chunk)
    return h.hexdigest()

def dump(path,obj):
    path.write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')

def verify_development_inputs():
    """Verify the locked, actually used development inputs without opening CTRP."""
    lock=json.loads((ROOT/'config/cohort_freeze_v1.lock.json').read_text(encoding='utf-8'))
    freeze_path=ROOT/'config/cohort_freeze_v1.json'
    if sha(freeze_path)!=lock['freeze_sha256']:raise ValueError('Cohort freeze no longer matches its lock')
    freeze=json.loads(freeze_path.read_text(encoding='utf-8'))
    recorded={str(Path(r['path'])):r['sha256'] for r in freeze['files'] if r.get('scope')!='original_read_only'}
    required=['results/development_sample_manifest.csv','results/main_drugs.csv','data_processed/expression_tpm.npy',
              'data_processed/expression_genes.tsv','data_processed/expression_samples.tsv',
              'data_processed/cell_line_dictionary.csv','data_manifest/original_source_sha256.csv','src/project_paths.py']
    for name in required:
        if str(Path(name)) not in recorded or sha(ROOT/name)!=recorded[str(Path(name))]:raise ValueError('Frozen development input changed: '+name)
    return {'freeze_sha256':lock['freeze_sha256'],'verified_development_files':required,'CTRP_response_files_opened':False}

def install_data_guard():
    """Fail before any pilot read of an external response or response index."""
    forbidden=[(ROOT/'data_raw/ctrp').resolve(),original_input('legacy_ctrp_response').resolve()]
    processed=[(ROOT/'data_processed'/n).resolve() for n in ['ctrp_curve_index.csv','ctrp_pair_availability.csv','ctrp_drug_availability.csv']]
    access=set();blocked=[]
    def hook(event,args):
        if event!='open' or not args or not isinstance(args[0],(str,bytes,os.PathLike)):return
        try:p=Path(os.fsdecode(args[0])).resolve()
        except (ValueError,OSError):return
        if p==forbidden[1] or p==forbidden[0] or p.is_relative_to(forbidden[0]) or p in processed:
            blocked.append(str(p));raise PermissionError('CTRP response access is forbidden in development pilot')
        if p.is_relative_to(ROOT/'data_processed') or p==original_input('gdsc2').resolve():access.add(str(p))
    sys.addaudithook(hook)
    return access,blocked

class ResourceMonitor:
    def __init__(self,path,interval):self.path=path;self.interval=interval;self.stop=threading.Event();self.rows=[];self.process=psutil.Process();self.result=None
    def start(self):
        self.started=time.perf_counter();self.cpu_start=self.process.cpu_times()
        def sample():
            while not self.stop.is_set():
                info=self.process.memory_info()
                self.rows.append({'elapsed_seconds':time.perf_counter()-self.started,'rss_bytes':info.rss,'windows_peak_working_set_bytes':getattr(info,'peak_wset',None)})
                self.stop.wait(self.interval)
        self.thread=threading.Thread(target=sample,daemon=True);self.thread.start();return self
    def close(self):
        if self.result is not None:return self.result
        self.stop.set();self.thread.join();info=self.process.memory_info()
        self.rows.append({'elapsed_seconds':time.perf_counter()-self.started,'rss_bytes':info.rss,'windows_peak_working_set_bytes':getattr(info,'peak_wset',None)})
        pd.DataFrame(self.rows).to_csv(self.path,index=False)
        cpu=self.process.cpu_times()
        self.result={'wall_seconds':time.perf_counter()-self.started,'cpu_seconds':cpu.user+cpu.system-self.cpu_start.user-self.cpu_start.system,
                'sampled_peak_rss_bytes':max(r['rss_bytes'] for r in self.rows),
                'windows_peak_working_set_bytes':getattr(info,'peak_wset',None),'memory_scope':'whole single pilot process including input preparation and libraries; not isolated model memory','sample_interval_seconds':self.interval}
        return self.result

def read_selected_responses(manifest,source,output):
    lookup={int(r['gdsc_source_row_number']):r for r in manifest.to_dict('records')}
    assert len(lookup)==len(manifest)
    records=[];wb=load_workbook(source,read_only=True,data_only=True);iterator=None
    try:
        iterator=wb.worksheets[0].iter_rows(values_only=True);header=next(iterator);col={k:i for i,k in enumerate(header)}
        for source_row,values in enumerate(iterator,2):
            if source_row not in lookup:continue
            r=lookup[source_row]
            assert str(values[col['DRUG_ID']])==str(r['gdsc_drug_id'])
            assert str(values[col['NLME_CURVE_ID']])==str(r['nlme_curve_id'])
            assert float(values[col['MIN_CONC']])==float(r['min_conc']) and float(values[col['MAX_CONC']])==float(r['max_conc'])
            y=float(values[col['AUC']]);assert np.isfinite(y) and 0<=y<=1
            records.append({**r,'response_auc':y})
    finally:
        if iterator is not None:iterator.close()
        wb.close()
    assert len(records)==len(lookup)
    response=pd.DataFrame(records).sort_values(['gdsc_drug_id','canonical_model_id'])
    response.to_csv(output,index=False,encoding='utf-8-sig');return response

def main():
    ACTIVE_RUN.clear()
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--config',type=Path,default=ROOT/'config/pilot_v0.1.json')
    ap.add_argument('--run-dir',type=Path,default=None)
    args=ap.parse_args();config=json.loads(args.config.read_text(encoding='utf-8'))
    out=args.run_dir or ROOT/'runs'/config['run_name']
    assert not out.exists(),f'Preserve previous run; choose a new run directory: {out}'
    out.mkdir(parents=True);(out/'features').mkdir()
    start=time.perf_counter();access,blocked=install_data_guard()
    monitor=ResourceMonitor(out/'resource_samples.csv',config['resources']['rss_sampling_interval_seconds']).start()
    ACTIVE_RUN.update({'out':out,'monitor':monitor,'config':config,'access':access,'blocked':blocked})
    packages=['numpy','pandas','scikit-learn','scipy','psutil','joblib','threadpoolctl','cloudpickle']
    run_meta={'config_sha256':sha(args.config),'cohort_lock':json.loads((ROOT/'config/cohort_freeze_v1.lock.json').read_text(encoding='utf-8')),
              'development_input_verification':verify_development_inputs(),
              'python':platform.python_version(),'platform':platform.platform(),'packages':{p:importlib.metadata.version(p) for p in packages},
              'logical_cpus':psutil.cpu_count(),'total_ram_bytes':psutil.virtual_memory().total,
              'configuration':config,'status':'running','external_response_accessed':False,
              'source_hashes':{str(p.relative_to(ROOT)):sha(p) for p in [args.config,ROOT/'src/run_development_pilot.py',ROOT/'src/modeling/features.py',ROOT/'src/modeling/nested_cv.py']}}
    dump(out/'run_manifest.json',run_meta)
    ACTIVE_RUN['metadata']=run_meta
    global_manifest=pd.read_csv(ROOT/'results/development_sample_manifest.csv',dtype=str,keep_default_na=False)
    selected=global_manifest[global_manifest.gdsc_drug_id.isin([str(d) for d in config['drug_ids']])].copy()
    assert set(selected.role)=={'development'}
    panel=pd.read_csv(ROOT/'results/main_drugs.csv',dtype=str)
    assert panel.iloc[:len(config['drug_ids'])].gdsc_drug_id.tolist()==[str(d) for d in config['drug_ids']]
    # Identity metadata only; no external response or expression distributions used.
    global_ids=global_manifest.drop_duplicates('canonical_model_id').sort_values('canonical_model_id').reset_index(drop=True)
    assert global_ids.biological_group_id.nunique()==len(global_ids)
    tissue_names=sorted(set(global_ids.tissue));tissue_code={t:i for i,t in enumerate(tissue_names)}
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        outer=make_grouped_splits(global_ids.biological_group_id,global_ids.tissue,config['outer_folds'],config['seed'])
        inner={oi:make_grouped_splits(global_ids.iloc[train].biological_group_id,global_ids.iloc[train].tissue,config['inner_folds'],config['seed']+100+oi) for oi,(train,test) in enumerate(outer)}
    split_warnings=[str(w.message) for w in caught]
    split_rows=[]
    for oi,(train,test) in enumerate(outer):
        for role,idx in [('train',train),('test',test)]:
            for n in idx:split_rows.append({**global_ids.iloc[n][['canonical_model_id','biological_group_id','tissue']].to_dict(),'outer_fold':oi,'inner_fold':-1,'role':role})
        for ii,(it,iv) in enumerate(inner[oi]):
            for role,idx in [('train',it),('validation',iv)]:
                for n in train[idx]:split_rows.append({**global_ids.iloc[n][['canonical_model_id','biological_group_id','tissue']].to_dict(),'outer_fold':oi,'inner_fold':ii,'role':role})
    pd.DataFrame(split_rows).to_csv(out/'global_fold_manifest.csv',index=False,encoding='utf-8-sig')
    source=original_input('gdsc2');expected=pd.read_csv(ROOT/'data_manifest/original_source_sha256.csv',dtype=str)
    expected=expected[expected.path.eq(str(source))];assert len(expected)==1 and sha(source)==expected.iloc[0].sha256
    response=read_selected_responses(selected,source,out/'development_responses.csv')
    genes=pd.read_csv(ROOT/'data_processed/expression_genes.tsv',sep='\t',dtype=str,keep_default_na=False)
    gene_idx=genes.index[genes.gene_type.eq('protein_coding')].to_numpy();stable=tuple(genes.iloc[gene_idx].ensembl_id)
    assert len(stable)==len(set(stable))
    matrix=np.load(ROOT/'data_processed/expression_tpm.npy',mmap_mode='r');assert matrix.dtype==np.float64
    preds=[];fold_metrics=[];choices=[];fit_events=[];timings=[];endpoint_audit=[];feature_benchmarks=[]
    id_pos={a:i for i,a in enumerate(global_ids.canonical_model_id)}
    fit_log=(out/'fit_events.jsonl').open('w',encoding='utf-8')
    def record(event):
        fit_events.append(event);fit_log.write(json.dumps(event,ensure_ascii=False,allow_nan=False)+'\n');fit_log.flush()
    try:
        with threadpool_limits(limits=config['resources']['blas_threads']):
            run_meta['runtime_threadpools']=threadpool_info()
            for drug in config['drug_ids']:
                data=response[response.gdsc_drug_id.eq(str(drug))].sort_values('canonical_model_id').reset_index(drop=True)
                rows=data.expression_row_index.astype(int).to_numpy()
                x=np.empty((len(data),len(gene_idx)+2),dtype=np.float64)
                x[:,:-2]=matrix[np.ix_(rows,gene_idx)];x[:,-2]=data.tissue.map(tissue_code);x[:,-1]=rows
                y=data.response_auc.to_numpy(dtype=np.float64);groups=data.biological_group_id.to_numpy()
                assert np.isfinite(x[:,:-2]).all() and (x[:,:-2]>=0).all()
                endpoint_audit.append({'gdsc_drug_id':drug,'n':len(y),'response_min':float(y.min()),'response_max':float(y.max()),'response_sd':float(y.std()),'n_unique_response':len(np.unique(y)),'dose_window':[float(data.min_conc.iloc[0]),float(data.max_conc.iloc[0])]})
                positions=np.array([id_pos[a] for a in data.canonical_model_id])
                for oi,(gtrain,gtest) in enumerate(outer):
                    train=np.flatnonzero(np.isin(positions,gtrain));test=np.flatnonzero(np.isin(positions,gtest))
                    # Intersect global inner roles with this drug while preserving local outer-train order.
                    outer_pos=positions[train];local_inner=[]
                    for it,iv in inner[oi]:
                        local_inner.append((np.flatnonzero(np.isin(outer_pos,gtrain[it])),np.flatnonzero(np.isin(outer_pos,gtrain[iv]))))
                    for kind in config['models']:
                        tick=time.perf_counter();context={'gdsc_drug_id':drug,'model':kind,'outer_fold':oi}
                        result=evaluate_outer_fold(x,y,groups,train,test,local_inner,kind,config,stable,audit_sink=lambda event,c=context:record({**c,**event}))
                        elapsed=time.perf_counter()-tick
                        fold_metrics.append({**context,**result['metrics'],'wall_seconds':elapsed,'unknown_tissue_samples':result['outer_unknown_tissue_samples'],'fallback_reason':result['outer_fit_audit']['fallback_reason']})
                        choices.append({**context,'chosen_params':result['chosen_params'],'best_inner_rmse':result['best_inner_rmse'],'grid_boundary':result['chosen_grid_boundary'],'inner_scores':result['inner_scores']})
                        if result['feature_audit'] is not None:dump(out/'features'/f'{drug}_{kind}_outer{oi}.json',result['feature_audit'])
                        timings.append({**context,'wall_seconds':elapsed,'fits':len(result['fit_events']),'fit_seconds_sum':sum(e['fit_seconds'] for e in result['fit_events'])})
                        for n,pred in zip(test,result['predictions']):
                            r=data.iloc[n];preds.append({**context,'canonical_model_id':r.canonical_model_id,'biological_group_id':r.biological_group_id,'expression_row_index':int(r.expression_row_index),'tissue':r.tissue,'observed_auc':float(y[n]),'predicted_auc':float(pred)})
                        pd.DataFrame(fold_metrics).to_csv(out/'fold_metrics.csv',index=False)
                        pd.DataFrame(preds).to_csv(out/'oof_predictions.partial.csv',index=False,encoding='utf-8-sig')
                        dump(out/'chosen_parameters.partial.json',choices)
                        print(json.dumps({'event':'outer_model_complete',**context,'n_train':len(train),'n_test':len(test),'wall_seconds':round(elapsed,2),'rmse':result['metrics']['rmse'],'spearman':result['metrics']['spearman']},ensure_ascii=False),flush=True)
                    if oi==0:
                        # Formal feature-budget timing only; no labels, tuning or performance used.
                        tick=time.perf_counter();b=FoldFeatureBuilder(kind='reo',n_genes=config['feature']['formal_reo_genes_transform_benchmark'],min_tpm=config['feature']['min_tpm'],min_detection_fraction=config['feature']['min_detection_fraction'],max_pair_tie_fraction=config['feature']['max_pair_tie_fraction'],stable_gene_ids=stable)
                        z=b.fit_transform(x[train,:-2]);elapsed=time.perf_counter()-tick
                        feature_benchmarks.append({'gdsc_drug_id':drug,'outer_fold':oi,'requested_genes':200,'actual_genes':len(b.gene_indices_),'retained_pairs':z.shape[1],'fit_transform_wall_seconds':elapsed,'matrix_bytes':z.nbytes,'model_fit_performed':False})
                        del z,b
                del x
    finally:fit_log.close()
    p=pd.DataFrame(preds);assert not p.duplicated(['gdsc_drug_id','model','canonical_model_id']).any()
    p.to_csv(out/'oof_predictions.csv',index=False,encoding='utf-8-sig')
    summary=[]
    for (drug,kind),sub in p.groupby(['gdsc_drug_id','model'],sort=False):
        m=metrics(sub.observed_auc,sub.predicted_auc)
        if kind=='mean':m['spearman']=None;m['spearman_unavailable_reason']='fold_constant_training_mean_not_an_individual_ranking_model'
        events=[e for e in fit_events if e['gdsc_drug_id']==drug and e['model']==kind]
        summary.append({'gdsc_drug_id':drug,'model':kind,**m,'fits':len(events),'convergence_warning_fits':sum(bool(e['warnings']) for e in events),'fallback_fits':sum(bool(e['fallback_reason']) for e in events),'wall_seconds':sum(t['wall_seconds'] for t in timings if t['gdsc_drug_id']==drug and t['model']==kind)})
    pd.DataFrame(summary).to_csv(out/'summary_metrics.csv',index=False,encoding='utf-8-sig')
    dump(out/'chosen_parameters.json',choices);dump(out/'feature200_benchmarks.json',feature_benchmarks);dump(out/'endpoint_audit.json',endpoint_audit)
    run_meta.update({'status':'completed','wall_seconds':time.perf_counter()-start,'resources':monitor.close(),'n_fits':len(fit_events),'n_oof_predictions':len(p),
                     'split_warnings':split_warnings,'threadpools':threadpool_info(),'input_files_accessed':sorted(access),'blocked_external_access_attempts':blocked,'external_response_accessed':False})
    dump(out/'run_manifest.json',run_meta)
    assert not blocked
    print(json.dumps({'event':'pilot_complete','run_dir':str(out),'n_fits':len(fit_events),'resources':run_meta['resources']},ensure_ascii=False),flush=True)

def run_with_failure_recording():
    """Preserve failure/resource evidence and completed outer checkpoints."""
    try:main()
    except BaseException as exc:
        if ACTIVE_RUN:
            meta=ACTIVE_RUN.get('metadata',{'configuration':ACTIVE_RUN['config']})
            meta.update({'status':'failed','failure_type':type(exc).__name__,'failure':str(exc),
                         'resources':ACTIVE_RUN['monitor'].close(),
                         'input_files_accessed':sorted(ACTIVE_RUN['access']),
                         'blocked_external_access_attempts':ACTIVE_RUN['blocked'],'external_response_accessed':False,
                         'partial_outputs_retained':True})
            dump(ACTIVE_RUN['out']/'run_manifest.json',meta)
        raise

if __name__=='__main__':run_with_failure_recording()
