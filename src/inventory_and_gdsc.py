"""Read original inputs without mutation; index identities and availability, not outcomes."""
from pathlib import Path
from datetime import datetime, timezone
from openpyxl import load_workbook
import csv, json, hashlib, math, sys
from project_paths import ORIGINAL_ROOT, original_input, original_label

ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ORIGINAL_ROOT
OLD = WORKSPACE / '3.Project' / 'Pan-Drugs-REOs'

def sha256(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda: f.read(8 * 1024 * 1024), b''): h.update(b)
    return h.hexdigest()

def write_csv(p, rows, fields=None):
    with p.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields or list(rows[0])); w.writeheader(); w.writerows(rows)

def main():
    files = []
    for p in WORKSPACE.rglob('*'):
        if not p.is_file() or ROOT in p.parents: continue
        s = p.stat()
        files.append({'relative_path': str(p.relative_to(WORKSPACE)), 'bytes': s.st_size,
                      'modified_utc': datetime.fromtimestamp(s.st_mtime, timezone.utc).isoformat()})
    write_csv(ROOT / 'data_manifest' / 'original_file_inventory.csv', files)
    chosen = [original_input(n) for n in ['gdsc2','legacy_cell_annotations','legacy_ctrp_bridge','legacy_ctrp_response','legacy_project_counts','ccle_tpm']]
    chosen += sorted((OLD/'scripts').glob('*.R'))
    manifest = [{'path': str(p.resolve()), 'relative_path': original_label(p),
                 'bytes': p.stat().st_size, 'sha256': sha256(p), 'access': 'read_only_original'} for p in chosen if p.is_file()]
    write_csv(ROOT / 'data_manifest' / 'original_source_sha256.csv', manifest)
    source = original_input('gdsc2')
    dest = ROOT / 'data_processed' / 'gdsc2_observation_index.csv'
    fields = {'DATASET':'dataset','NLME_RESULT_ID':'nlme_result_id','NLME_CURVE_ID':'nlme_curve_id',
              'COSMIC_ID':'cosmic_id','CELL_LINE_NAME':'cell_line_name','SANGER_MODEL_ID':'sanger_model_id',
              'TCGA_DESC':'tcga_desc','DRUG_ID':'gdsc_drug_id','DRUG_NAME':'gdsc_drug_name',
              'PUTATIVE_TARGET':'putative_target','PATHWAY_NAME':'pathway_name','WEBRELEASE':'webrelease',
              'MIN_CONC':'min_conc','MAX_CONC':'max_conc'}
    wb = load_workbook(source, read_only=True, data_only=True)
    assert len(wb.worksheets) == 1
    it = wb.worksheets[0].iter_rows(values_only=True)
    header = next(it)
    n = missing = 0
    with dest.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['source_row_number', *fields.values(), 'response_available'])
        writer.writeheader()
        for i, values in enumerate(it, 2):
            row = dict(zip(header, values))
            if row.get('DRUG_ID') is None: continue
            auc = row.get('AUC')
            available = isinstance(auc, (int, float)) and math.isfinite(auc)
            rec = {'source_row_number': i, **{v: row.get(k) for k,v in fields.items()}, 'response_available': int(available)}
            writer.writerow(rec); n += 1; missing += not available
    wb.close()
    summary = {'original_file_count': len(files), 'hashed_input_and_script_count': len(manifest),
               'gdsc2_rows': n, 'response_missing_or_nonfinite': missing,
               'response_magnitudes_exported_or_summarized': False,
               'original_source': str(source.resolve()), 'source_sha256': sha256(source)}
    (ROOT/'results'/'initial_inventory_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False))

if __name__ == '__main__': main()
