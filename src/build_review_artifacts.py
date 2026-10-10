"""Create human-readable cohort/download review artifacts; never read response values."""
from pathlib import Path
import csv,json,platform,importlib.metadata
import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
DOWNLOADS=[
('CCLE RSEM TPM 20180929','本地已具备，完成结构与数值检查','核心表达输入','无需重新下载；未核对远端压缩包字节同一性','1.数据/2.CCLE-耐药敏感表达数据/0.CCLE原始数据/1.CCLE_RNAseq_rsem_genes_tpm_20180929.txt','https://data.broadinstitute.org/ccle/CCLE_RNAseq_rsem_genes_tpm_20180929.txt.gz'),
('GDSC2 fitted response 27Oct23','本地已具备，已建逐行索引','开发响应来源','无需重新下载；模型阶段才读取开发AUC','3.Project/Pan-Drugs-REOs/data/GDSC2_fitted_dose_response_27Oct23.xlsx','https://www.cancerrxgene.org/downloads/bulk_download'),
('DepMap 24Q4 Public v1 Model + README','已下载，官方MD5核验通过','ACH稳定ID及同组织样本来源关系','完成','data_raw/cell_metadata','https://doi.org/10.25452/figshare.plus.27993248.v1'),
('Sanger Cell Model Passports API快照','已下载，2292模型及关系','SIDM/COSMIC/ACH交叉映射和patient/parent关系','完成；记录2026-10-10快照','data_raw/sanger_metadata','https://depmap.sanger.ac.uk/documentation/api/'),
('GDSC release8.5 compound annotations','已下载，覆盖295个本地药物ID','官方药物名称及同义名','完成；原文件不包含实测化学结构','data_raw/drug_metadata/screened_compounds_rel_8.5.csv','https://cog.sanger.ac.uk/cancerrxgene/GDSC_release8.5/screened_compounds_rel_8.5.csv'),
('CTRPv2 ExpandedDataset 曲线和元数据','已下载所需9个成员并核对CRC','恢复compound/cell/experiment ID和SNP身份质控','完成；未下载整个342MB包/全部原始孔级数据','data_raw/ctrp','https://d1hp0ufb0xxisr.cloudfront.net/CTD2/Broad/CTRPv2.0_2015_ctd2_ExpandedDataset.zip'),
('PubChem compound properties','已缓存80个名称查询及结构','对照CTRP SMILES生成的完整InChIKey','完成；名称解析结构并非对GDSC药瓶的独立鉴定','data_raw/drug_metadata','https://pubchem.ncbi.nlm.nih.gov/docs/pug-rest'),
('匹配CCLE版本的GENCODE19注释','已下载','基因类型和gene ID注释','完成；1582未精确匹配基因保留，不能当作非编码','data_raw/expression/gencode.v19.genes.v7_model.patched_contigs.gtf.gz','https://data.broadinstitute.org/ccle/gencode.v19.genes.v7_model.patched_contigs.gtf.gz'),
('CTRP/GDSC共同浓度窗的曲线重积分材料','条件性补充，尚未完成可行性核查','增强跨库终点可比性','现有CTRP拟合参数可先评估；仅在需要重拟合时补孔级剂量数据；开发pilot之前写入后续分析规范','待定',''),
('独立来源表达矩阵','可选，未下载/未选定','若论文需要跨表达平台泛化证据','需核验来源确实独立于当前CCLE RNA-seq后再选；现阶段不可称跨表达平台验证','待定',''),
('PRISM二级剂量响应','可选，未下载/未做覆盖核查','第三筛选库的补充验证','只在主流程稳定且时间允许时；避免扩大一个月工作量','待定',''),
('GEO/单细胞元数据与原始计数','可选，暂不重新下载','原计划的探索性单细胞补充','主分析完成后再审查独立重复和用药结局；不得替代当前外验','待定',''),
('全量FASTQ、TCGA生存、额外多组学','本阶段不需要','不属于当前研究问题的必要输入','不下载','不适用','')]

def main():
    register=ROOT/'data_manifest/download_register.csv'
    with register.open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['dataset','status','purpose','action_or_limit','path','source_url']);w.writerows(DOWNLOADS)
    env={'python':platform.python_version(),'platform':platform.platform(),'packages':{p:importlib.metadata.version(p) for p in ['numpy','pandas','openpyxl','rdkit']}}
    (ROOT/'config/curation_environment.json').write_text(json.dumps(env,indent=2),encoding='utf-8')
    wb=Workbook();intro=wb.active;intro.title='阅读说明'
    for row in [
        ['数据核验与队列锁定','阶段1；目标期刊BMC Genomics'],
        ['时间','2026-10-10'],['开发/外验唯一模型','627 / 180；药物间重复使用模型，不能累加为独立样本'],
        ['主分析','8种预先按身份和样本覆盖筛选的化合物；不是8种均获批临床药物'],
        ['外验边界','已知身份和同源组与全部本地GDSC2不重叠；CTRP在历史研究中用过'],
        ['终点边界','当前支持跨筛选实验的排序迁移；共同CCLE表达不能称跨表达平台验证'],
        ['尚未开展','模型拟合、特征筛选、外验性能计算'],
        ['权威机器输入','CSV、config/cohort_freeze_v1.json及对应.lock.json；此工作簿仅供查阅'],
        ['质控','reports/validation_summary.json'],
        ['原文件','只读；新输出全部位于rebuild_v2']]: intro.append(row)
    sheetfiles=[('8药物面板','results/main_drugs.csv'),('295药物样本量','results/all_295_GDSC_drug_sample_counts.csv'),
                ('结构通过候选','results/all_structure_supported_drug_coverage.csv'),('细胞系总映射','data_processed/cell_line_dictionary.csv'),
                ('GDSC细胞映射','data_processed/gdsc2_cell_mapping.csv'),('CTRP细胞映射','data_processed/ctrp_cell_mapping.csv'),
                ('药物身份冲突','data_processed/drug_identity_ambiguous.csv'),('浓度窗决策','results/gdsc_dose_range_decisions.csv'),
                ('开发组织分布','results/development_tissue_counts.csv'),('外验组织分布','results/external_primary_tissue_counts.csv'),
                ('下载清单','data_manifest/download_register.csv')]
    for title,path in sheetfiles:
        ws=wb.create_sheet(title)
        with (ROOT/path).open(encoding='utf-8-sig',newline='') as f:
            for row in csv.reader(f): ws.append(row)
        ws.auto_filter.ref=ws.dimensions
    for ws in wb:
        ws.freeze_panes='B2';ws.sheet_view.showGridLines=False
        for c in ws[1]: c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='204A63');c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[1].height=38
        for cells in ws.iter_rows(min_row=2):
            for c in cells:
                c.alignment=Alignment(vertical='top')
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F5F8')
                if c.data_type=='f': c.data_type='s'  # metadata must never execute as an Excel formula
        for col in range(1,ws.max_column+1):
            length=max(len(str(ws.cell(row,col).value or '')) for row in range(1,min(ws.max_row,80)+1))
            ws.column_dimensions[get_column_letter(col)].width=min(48,max(15,length+2))
    dest=ROOT/'reports/阶段1_队列与下载清单.xlsx';wb.save(dest)
    check=load_workbook(dest,read_only=True);assert len(check.sheetnames)==12 and check['8药物面板'].max_row==9;check.close()
    print(json.dumps({'workbook':str(dest),'sheets':len(wb.sheetnames),'download_items':len(DOWNLOADS)},ensure_ascii=False))

if __name__=='__main__':main()
