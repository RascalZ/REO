"""Build local bilingual submission v2 and title-consistent cover; no study data."""
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / 'docs/scope_diagnostics_revision_v1'
OLD = ROOT / 'manuscript_v1/submission_final_v1'
NEW = ROOT / 'manuscript_v1/submission_final_v2'
ZH = ROOT / 'manuscript_v1/submission_final_v2_zh'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    protected = json.loads((AUDIT/'local_private/preservation_manifest.json').read_text(encoding='utf-8'))['files']
    assert all(sha(ROOT/p) == h for p, h in protected.items())
    old_md = (OLD/'Manuscript_BMC_Genomics_submission.md').read_text(encoding='utf-8')
    new_md = (NEW/'Manuscript_BMC_Genomics_submission_v2.md').read_text(encoding='utf-8')
    old_title = next(x.split('**Title:** ', 1)[1] for x in old_md.splitlines() if x.startswith('**Title:** '))
    new_title = next(x.split('**Title:** ', 1)[1] for x in new_md.splitlines() if x.startswith('**Title:** '))
    figure = ROOT/'figures/manuscript_scope_diagnostics_v1/Figure4_loto_and_external.png'
    assert figure.is_file()
    shutil.copy2(figure, NEW/'Figures'/figure.name)
    for folder in ['Figures', 'Additional_files', 'Supporting_sources']:
        for p in (OLD/folder).iterdir():
            if p.name == figure.name:
                continue
            assert sha(p) == sha(NEW/folder/p.name), p
    cover = Document(OLD/'Cover_letter_BMC_Genomics.docx')
    edits = 0
    before = 'The study does not establish an average pair-encoding advantage;'
    after = 'The evaluated complete-pair elastic-net pipelines did not establish an average predictive advantage;'
    for p in cover.paragraphs:
        new = p.text.replace(old_title, new_title).replace('27 October 2026', '29 October 2026').replace(before, after)
        new = new.replace('its contribution is an auditable assessment', "the study's contribution is an auditable assessment")
        if new != p.text:
            # These cover paragraphs are uniform-format runs in the source.
            if p.runs:
                p.runs[0].text = new
                for run in p.runs[1:]:
                    run.text = ''
            else:
                p.add_run(new)
            edits += 1
    assert edits == 4, edits
    cover.save(NEW/'Cover_letter_BMC_Genomics.docx')
    (NEW/'Cover_letter_BMC_Genomics.txt').write_text('\n\n'.join(p.text for p in cover.paragraphs)+'\n', encoding='utf-8')
    node = os.environ.get('REO_DOC_NODE', 'node')
    for builder, extra in [('build_submission_scope_revision_v2.js', ['--embed-figures']),
                           ('build_submission_scope_revision_v2_zh.js', [])]:
        subprocess.run([node, str(ROOT/'manuscript_v1'/builder), *extra], cwd=ROOT, check=True)
    assert all(sha(ROOT/p) == h for p, h in protected.items())
    record = {'date':'2026-10-29', 'preserved_source_files':len(protected),
              'cover_changed_paragraphs':edits, 'cover_scope':'Title, date, evaluated-pipeline wording only',
              'unchanged_additional_files':14, 'unchanged_figures':[1,2,3],
              'new_figure4_source_sha256':sha(figure), 'training_or_scoring_performed':False,
              'individual_CTRP_access':False}
    (AUDIT/'document_build.json').write_text(json.dumps(record, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(record))


if __name__ == '__main__':
    main()
