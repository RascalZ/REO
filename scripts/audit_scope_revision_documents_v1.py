"""Audit and render editorial artifacts; reads no model, response, or OOF data."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
from zipfile import ZipFile
import pdfplumber
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT/'docs/scope_diagnostics_revision_v1'
PRIVATE = AUDIT/'local_private'
OUT = ROOT/'manuscript_v1/submission_final_v2'
ZH = ROOT/'manuscript_v1/submission_final_v2_zh'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(s):
    return re.sub(r'\s+', '', s.replace('**', '').replace('*', '').replace('`', ''))


def main():
    protected = json.loads((PRIVATE/'preservation_manifest.json').read_text(encoding='utf-8'))['files']
    assert all(sha(ROOT/p) == h for p,h in protected.items())
    old = ROOT/'manuscript_v1/submission_final_v1'
    report = {'date':'2026-10-29', 'scope':'Editorial and aggregate-display audit; not model replication',
              'preserved_source_files':len(protected), 'all_protected_sources_unchanged':True,
              'training_or_scoring_performed':False, 'individual_CTRP_access':False,
              'documents':{}, 'unchanged_support_files':[]}
    for folder in ['Additional_files','Supporting_sources']:
        for p in (old/folder).iterdir():
            assert sha(p)==sha(OUT/folder/p.name), p
            report['unchanged_support_files'].append(str(p.relative_to(old)))
    for n,name in enumerate(['Figure1_study_design.png','Figure2_development_primary.png','Figure3_representation_and_zero_solutions.png'],1):
        assert sha(old/'Figures'/name)==sha(OUT/'Figures'/name)
    en = (OUT/'Manuscript_BMC_Genomics_submission_v2.md').read_text(encoding='utf-8')
    old_en = (old/'Manuscript_BMC_Genomics_submission.md').read_text(encoding='utf-8')
    def section(text,name):
        return text.split('## '+name+'\n',1)[1].split('\n## ',1)[0].strip()
    assert section(en,'References')==section(old_en,'References')
    report['references_byte_identical']=True
    report['abstract_word_count']=len(section(en,'Abstract').split())
    assert report['abstract_word_count']<=350
    for title in ['Ethics approval and consent to participate','Consent for publication','Competing interests','Funding',"Authors' contributions",'Acknowledgements']:
        def sub(text):
            return text.split('### '+title+'\n',1)[1].split('\n### ',1)[0].split('\n## ',1)[0].strip()
        assert sub(en)==sub(old_en), title
    report['author_declaration_facts_unchanged']=True
    for kind,base in [('english',OUT/'Manuscript_BMC_Genomics_submission_v2'),
                      ('chinese',ZH/'Manuscript_BMC_Genomics_submission_v2_zh'),
                      ('cover',OUT/'Cover_letter_BMC_Genomics')]:
        docx,pdf=base.with_suffix('.docx'),base.with_suffix('.pdf')
        qa=PRIVATE/'render'/kind; qa.mkdir(parents=True,exist_ok=True)
        d={'docx_sha256':sha(docx),'pdf_sha256':sha(pdf),'text_outside_page':0,
           'suspect_missing_glyphs':0,'figure_pages':[], 'font_embedding':{}}
        document=Document(docx)
        if kind!='cover':
            md=base.with_suffix('.md').read_text(encoding='utf-8')
            body=''.join(p.text for p in document.paragraphs)
            tables=''.join(c.text for t in document.tables for row in t.rows for c in row.cells)
            normalized=norm(body+tables)
            absent=[]
            for i,line in enumerate(md.splitlines(),1):
                line=line.strip()
                if not line or line.startswith('#') or line.startswith('|'):
                    continue
                if norm(line) not in normalized:
                    absent.append(i)
            d['markdown_lines_missing_from_docx']=absent
            assert not absent,(kind,absent)
            assert len(document.tables)==2
            d['native_tables']=2
            with ZipFile(docx) as z:
                image_hashes=[hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]
            desired=[sha(p) for p in (OUT/'Figures').glob('*.png')]
            assert sorted(image_hashes)==sorted(desired)
            d['embedded_figure_hashes_match_package']=True
        reader=PdfReader(pdf); d['pages']=len(reader.pages)
        for page in reader.pages:
            for name,font in page.get('/Resources',{}).get('/Font',{}).items():
                font=font.get_object(); descendants=font.get('/DescendantFonts')
                underlying=descendants[0].get_object() if descendants else font
                desc=underlying.get('/FontDescriptor')
                embedded=bool(desc and any(k in desc.get_object() for k in ['/FontFile','/FontFile2','/FontFile3']))
                d['font_embedding'][str(font.get('/BaseFont',name))]=embedded
        with pdfplumber.open(pdf) as source:
            texts=[]
            for n,page in enumerate(source.pages,1):
                texts.append(page.extract_text(layout=False) or '')
                d['text_outside_page']+=sum(c['x0'] < -.5 or c['x1']>page.width+.5 or c['top']<-.5 or c['bottom']>page.height+.5 for c in page.chars)
                d['suspect_missing_glyphs']+=sum(c['text']=='\ufffd' for c in page.chars)
                if page.images:d['figure_pages'].append(n)
                assert len(texts[-1].strip())>20,(kind,n,'blank page')
            (qa/'extracted_text.txt').write_text('\n\f\n'.join(texts),encoding='utf-8')
        assert d['text_outside_page']==0 and d['suspect_missing_glyphs']==0
        assert all(d['font_embedding'].values())
        if kind!='cover':assert len(d['figure_pages'])==4
        if kind=='english':
            d['sequential_page_numbers']=all(re.search(r'Page\s+'+str(i)+r'\s*$',text) for i,text in enumerate(texts,1))
        elif kind=='chinese':
            d['sequential_page_numbers']=all(re.search(r'第\s*'+str(i)+r'\s*页',text) for i,text in enumerate(texts,1))
        else:d['sequential_page_numbers']=None
        if kind!='cover':assert d['sequential_page_numbers']
        pages=qa/'pages'; pages.mkdir(exist_ok=True)
        subprocess.run(['pdftoppm','-scale-to','2000','-png',str(pdf),str(pages/'page')],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
        images=sorted(pages.glob('page-*.png')); assert len(images)==d['pages']
        sheets=[]
        for start in range(0,len(images),12):
            sheet=Image.new('RGB',(1800,1760),'#d6d6d6'); draw=ImageDraw.Draw(sheet)
            for slot,p in enumerate(images[start:start+12]):
                im=Image.open(p).convert('RGB'); im.thumbnail((440,550))
                x=(slot%4)*450+(450-im.width)//2; y=(slot//4)*585+27
                sheet.paste(im,(x,y)); draw.text(((slot%4)*450+8,(slot//4)*585+5),f'{kind} page {start+slot+1}',fill='black')
            path=qa/f'contact_{start//12+1}.png';sheet.save(path);sheets.append(str(path.relative_to(ROOT)))
        d['contact_sheets']=sheets; d['full_page_directory']=str(pages.relative_to(ROOT))
        report['documents'][kind]=d
    report['mechanical_checks_passed']=True
    (AUDIT/'mechanical_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'pass':True,'documents':{k:{'pages':v['pages'],'figure_pages':v['figure_pages']} for k,v in report['documents'].items()}},ensure_ascii=False))


if __name__=='__main__':
    main()
